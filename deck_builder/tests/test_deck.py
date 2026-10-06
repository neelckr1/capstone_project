"""
TDD Test Suite for AegisMAS Capstone Presentation Deck
Validates content, styling, structure, and rendering at every step.
"""
import unittest
import os
import pptx
from pptx.dml.color import RGBColor
from deck_builder.config import (
    TEMPLATE_PATH, OUTPUT_PPTX_PATH, RENDERS_DIR
)
from deck_builder.builder import build_presentation
from deck_builder.render import render_presentation_to_images

class TestAegisMASDeck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """Build the presentation once for testing."""
        cls.prs = build_presentation()
        cls.slides = list(cls.prs.slides)

    def test_slide_count(self):
        """Verify the presentation has exactly 10 slides."""
        self.assertEqual(len(self.slides), 10, "Presentation must contain exactly 10 slides.")

    def test_slide_1_title(self):
        """Verify Slide 1 title, metadata, sub-theme, and banner."""
        s1 = self.slides[0]
        full_text = " ".join([shape.text_frame.text for shape in s1.shapes if shape.has_text_frame])
        
        self.assertIn("AegisMAS", full_text)
        self.assertIn("Mind Viruses", full_text)
        self.assertIn("UE24CS320A-CAP-2026", full_text)
        self.assertIn("PES University", full_text)
        self.assertIn("Non-Autoregressive Triaging", full_text)
        self.assertIn("Defending the Autonomous Frontier", full_text)

    def test_slide_2_outline(self):
        """Verify Slide 2 outline has 7 numbered agenda items and callout."""
        s2 = self.slides[1]
        title_shape = [s for s in s2.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Presentation Outline")
        
        full_text = " ".join([shape.text_frame.text for shape in s2.shapes if shape.has_text_frame])
        for num in ["01 ›", "02 ›", "03 ›", "04 ›", "05 ›", "06 ›", "07 ›"]:
            self.assertIn(num, full_text, f"Missing outline item {num}")
        self.assertIn("Option 3: Dual-Anchor Hybrid", full_text)

    def test_slide_3_problem_statement(self):
        """Verify Slide 3 problem statement and R0 callout cards."""
        s3 = self.slides[2]
        title_shape = [s for s in s3.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Problem Statement")
        
        full_text = " ".join([shape.text_frame.text for shape in s3.shapes if shape.has_text_frame])
        self.assertIn("Inter-Agent Trust Asymmetry", full_text)
        self.assertIn("Mind Virus", full_text)
        self.assertIn("Llama Guard", full_text)
        self.assertIn("Monolithic LLM", full_text)
        self.assertIn("Formal Problem Statement", full_text)
        self.assertIn("R₀ > 1", full_text)
        self.assertIn("R₀ < 1", full_text)

    def test_slide_4_scope_and_feasibility(self):
        """Verify Slide 4 scope, challenges, and State-Drift Vector mitigation."""
        s4 = self.slides[3]
        title_shape = [s for s in s4.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Scope and Feasibility Study")
        
        full_text = " ".join([shape.text_frame.text for shape in s4.shapes if shape.has_text_frame])
        self.assertIn("In-Scope Deliverables", full_text)
        self.assertIn("Out-of-Scope", full_text)
        self.assertIn("State-Drift Vector", full_text)
        self.assertIn("Guardrail Curse", full_text)
        self.assertIn("AgentBench / GAIA", full_text)
        self.assertIn("Key Insight", full_text)

    def test_slide_5_background_work(self):
        """Verify Slide 5 paradigm shift, 4 research citations, and tags."""
        s5 = self.slides[4]
        title_shape = [s for s in s5.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Background Work & Domain Context")
        
        full_text = " ".join([shape.text_frame.text for shape in s5.shapes if shape.has_text_frame])
        self.assertIn("Paradigm Shift", full_text)
        self.assertIn("Anthropic", full_text)
        self.assertIn("OpenAI–Hugging Face Incident", full_text)
        self.assertIn("AutoInject", full_text)
        self.assertIn("CivicShield", full_text)

        # Check stale TextBox 2 is empty
        tb2 = [s for s in s5.shapes if s.name == "TextBox 2"]
        if tb2:
            self.assertEqual(tb2[0].text_frame.text.strip(), "", "Stale template box must be empty.")

    def test_slide_6_applications(self):
        """Verify Slide 6 enterprise applications and economic justification."""
        s6 = self.slides[5]
        title_shape = [s for s in s6.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Real-World Applications & Use Cases")
        
        full_text = " ".join([shape.text_frame.text for shape in s6.shapes if shape.has_text_frame])
        self.assertIn("Enterprise Agent Marketplaces", full_text)
        self.assertIn("Autonomous DevOps", full_text)
        self.assertIn("Financial & Government Meshes", full_text)
        self.assertIn("Economic Justification", full_text)

    def test_slide_7_deliverables(self):
        """Verify Slide 7 deliverables across Capstone I, II, III."""
        s7 = self.slides[6]
        title_shape = [s for s in s7.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Expected Deliverables by Phase")
        
        full_text = " ".join([shape.text_frame.text for shape in s7.shapes if shape.has_text_frame])
        self.assertIn("Capstone-I", full_text)
        self.assertIn("Capstone-II", full_text)
        self.assertIn("Capstone-III", full_text)

    def test_slide_8_timeline_gantt_table(self):
        """Verify Slide 8 Gantt table geometry and team responsibility rows."""
        s8 = self.slides[7]
        title_shape = [s for s in s8.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "16-Week Project Timeline & Team Effort")
        
        tables = [s.table for s in s8.shapes if s.has_table]
        self.assertEqual(len(tables), 1, "Slide 8 must contain exactly 1 Gantt table.")
        table = tables[0]
        self.assertEqual(len(table.rows), 10, "Table must have 10 rows.")
        self.assertEqual(len(table.columns), 9, "Table must have 9 columns.")
        
        # Verify header
        self.assertIn("Task / Milestone", table.cell(0, 0).text)
        self.assertIn("Wk 1-2", table.cell(0, 1).text)
        self.assertIn("Wk 15-16", table.cell(0, 8).text)
        
        # Verify team responsibility row
        self.assertIn("Team Responsibility", table.cell(9, 0).text)
        self.assertIn("Member 1", table.cell(9, 1).text)
        self.assertIn("Member 2", table.cell(9, 4).text)
        self.assertIn("Member 3", table.cell(9, 7).text)

    def test_slide_9_methodology_and_diagram(self):
        """Verify Slide 9 architecture explanation and visual flow nodes."""
        s9 = self.slides[8]
        title_shape = [s for s in s9.shapes if s.name == "Text Box 34"][0]
        self.assertEqual(title_shape.text_frame.text.strip(), "Proposed Methodology: AegisMAS Defense-in-Depth")
        
        full_text = " ".join([shape.text_frame.text for shape in s9.shapes if shape.has_text_frame])
        self.assertIn("System-1 Non-Autoregressive Decision Gate", full_text)
        self.assertIn("System-2 Deliberative Trajectory Reasoning", full_text)
        self.assertIn("Digital Twin", full_text)
        self.assertIn("Green", full_text)
        self.assertIn("Amber", full_text)
        self.assertIn("Red", full_text)

    def test_slide_10_conclusion(self):
        """Verify Slide 10 conclusion card and 3 metric highlight boxes."""
        s10 = self.slides[9]
        full_text = " ".join([shape.text_frame.text for shape in s10.shapes if shape.has_text_frame])
        self.assertIn("Thank You", full_text)
        self.assertIn("AegisMAS — Defending the Autonomous Frontier", full_text)
        self.assertIn("R₀ > 1  ➔  R₀ < 1", full_text)
        self.assertIn("<50ms  Triage", full_text)
        self.assertIn(">75%  Cost Savings", full_text)
        self.assertIn("PES University", full_text)

    def test_no_template_placeholders_remaining(self):
        """Assert that no unedited placeholder prompts from the original template remain."""
        stale_prompts = [
            "Well defined problem statement",
            "It should clearly specify the problem",
            "Provide an overview of scope it entails",
            "Possible Shortcomings/Challenges",
            "Describe the applications and use cases of your project",
            "Provide any other information you wish to add on"
        ]
        for idx, slide in enumerate(self.slides, start=1):
            slide_text = " ".join([s.text_frame.text for s in slide.shapes if s.has_text_frame])
            for prompt in stale_prompts:
                self.assertNotIn(prompt, slide_text, f"Stale template prompt '{prompt}' found on Slide {idx}")

    def test_rendering_pipeline(self):
        """Test that Keynote + PyMuPDF pipeline exports and renders all 10 slide images."""
        image_paths = render_presentation_to_images(OUTPUT_PPTX_PATH, RENDERS_DIR, dpi=150)
        self.assertEqual(len(image_paths), 10, "Expected 10 rendered slide images.")
        for img_path in image_paths:
            self.assertTrue(os.path.exists(img_path), f"Rendered image does not exist: {img_path}")
            size_bytes = os.path.getsize(img_path)
            self.assertGreater(size_bytes, 50000, f"Rendered slide {img_path} is suspiciously small ({size_bytes} bytes).")

if __name__ == "__main__":
    unittest.main()
