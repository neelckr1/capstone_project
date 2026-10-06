"""
Deck Builder Orchestrator: loads template, applies slide builders, and saves PPTX.
"""
import pptx
import shutil
from deck_builder.config import (
    TEMPLATE_PATH, OUTPUT_PPTX_PATH, OUTPUT_PPTX_ROOT_PATH
)
from deck_builder.slides.slide_01_title import build_slide_1
from deck_builder.slides.slide_02_outline import build_slide_2
from deck_builder.slides.slide_03_problem import build_slide_3
from deck_builder.slides.slide_04_scope import build_slide_4
from deck_builder.slides.slide_05_background import build_slide_5
from deck_builder.slides.slide_06_applications import build_slide_6
from deck_builder.slides.slide_07_deliverables import build_slide_7
from deck_builder.slides.slide_08_timeline import build_slide_8
from deck_builder.slides.slide_09_methodology import build_slide_9
from deck_builder.slides.slide_10_conclusion import build_slide_10

SLIDE_BUILDERS = [
    build_slide_1,
    build_slide_2,
    build_slide_3,
    build_slide_4,
    build_slide_5,
    build_slide_6,
    build_slide_7,
    build_slide_8,
    build_slide_9,
    build_slide_10,
]

def build_presentation(template_path=TEMPLATE_PATH, output_path=OUTPUT_PPTX_PATH):
    """Build the entire presentation by running all slide builders sequentially."""
    prs = pptx.Presentation(template_path)
    slides = list(prs.slides)
    
    if len(slides) != len(SLIDE_BUILDERS):
        raise ValueError(f"Template has {len(slides)} slides, but expected {len(SLIDE_BUILDERS)}.")

    for idx, (slide, builder) in enumerate(zip(slides, SLIDE_BUILDERS), start=1):
        builder(slide)

    prs.save(output_path)
    if output_path != OUTPUT_PPTX_ROOT_PATH:
        prs.save(OUTPUT_PPTX_ROOT_PATH)
    
    return prs
