"""
Slide 1: Title Slide Builder
"""
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN
from deck_builder.config import (
    C_INDIGO, C_DARK, C_MUTED, C_WHITE
)
from deck_builder.helpers import (
    clear_tf, write_header_bullet, add_para, add_run, add_textbox
)

def build_slide_1(slide):
    for shape in slide.shapes:
        if shape.name == "Google Shape;26;p3":
            tf = clear_tf(shape)
            write_header_bullet(
                tf, "Project Title   :",
                "AegisMAS: Active Defense-in-Depth Against Self-Propagating Mind Viruses in Multi-Agent Systems",
                label_color=C_INDIGO, body_color=C_DARK, size=15, first=True, space_after=6
            )
            write_header_bullet(
                tf, "Project ID       :", "UE24CS320A-CAP-2026",
                label_color=C_INDIGO, body_color=C_DARK, size=15, space_after=6
            )
            write_header_bullet(
                tf, "Project Guide :", "[Faculty Guide Name], Dept. of CSE, PES University",
                label_color=C_INDIGO, body_color=C_DARK, size=15, space_after=6
            )
            write_header_bullet(
                tf, "Project Team  :", "Neel Chandrakar & Team [USN: 01FE... / PES University]",
                label_color=C_INDIGO, body_color=C_DARK, size=15, space_after=8
            )
            p_sub = add_para(tf, level=0, space_after=4)
            add_run(
                p_sub,
                "Sub-theme: Real-Time Non-Autoregressive Triaging & Co-Evolutionary Red-Teaming for MAS Security",
                bold=False, italic=True, size=12, color=C_MUTED
            )

    # Add bottom decorative banner on Slide 1
    add_textbox(
        slide, Emu(1371600), Emu(5800000), Emu(9448800), Emu(500000),
        "Defending the Autonomous Frontier  •  PES University Capstone 2025–26",
        size=13, bold=True, color=C_WHITE, bg_color=C_INDIGO, align=PP_ALIGN.CENTER
    )
