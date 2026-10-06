"""
Slide 8: Project Timeline & Team Effort Builder
"""
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN
from deck_builder.config import (
    C_INDIGO, C_DARK, C_CYAN, C_MUTED, C_WHITE,
    C_CALLOUT_BG, C_CRIMSON, C_AMBER, C_GREEN, FONT_FAMILY
)
from deck_builder.helpers import (
    clear_tf, set_slide_title
)

def build_slide_8(slide):
    set_slide_title(slide, "Text Box 34", "Capstone (Phase-I & Phase-II) Project Timeline")
    for shape in slide.shapes:
        if shape.name == "TextBox 4":
            tf = clear_tf(shape)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "(Structured Gantt chart and effort allocation detailed below)"
            run.font.name = FONT_FAMILY
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = C_MUTED

    # Add Gantt Table
    t_left   = Emu(1066800)
    t_top    = Emu(2000000)
    t_width  = Emu(10600000)
    t_height = Emu(4200000)
    table_shape = slide.shapes.add_table(10, 9, t_left, t_top, t_width, t_height)
    table = table_shape.table

    # Set Column Widths
    table.columns[0].width = Emu(3800000)  # Task column
    for col_idx in range(1, 9):
        table.columns[col_idx].width = Emu(850000)

    headers = ["Task / Milestone", "Wk 1-2", "Wk 3-4", "Wk 5-6", "Wk 7-8", "Wk 9-10", "Wk 11-12", "Wk 13-14", "Wk 15-16"]
    for col_idx, htext in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_INDIGO
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT
        run = p.add_run()
        run.text = htext
        run.font.name = FONT_FAMILY
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = C_WHITE

    tasks = [
        ("Literature Review & Threat Modeling", [1, 2]),
        ("MAS Testbed Setup (Sequential + Marketplace)", [2, 3]),
        ("Red-Agent Fuzzer & Baseline R₀ Measurement", [3, 4]),
        ("State-Drift Vector Encoder", [4, 5]),
        ("System-1 (Clef/Laya) Integration & Latency Benchmarking", [5, 6]),
        ("System-2 Deliberation Engine + HITL Dashboard", [6, 7]),
        ("Comprehensive Benchmarking & Threshold Tuning", [7, 8]),
        ("Paper Draft, Documentation & Presentation", [8])
    ]

    for row_idx, (tname, active_cols) in enumerate(tasks, start=1):
        c0 = table.cell(row_idx, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = C_WHITE
        c0.text = ""
        p = c0.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = tname
        run.font.name = FONT_FAMILY
        run.font.size = Pt(10)
        run.font.color.rgb = C_DARK
        
        for c_idx in range(1, 9):
            cell = table.cell(row_idx, c_idx)
            cell.fill.solid()
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.font.name = FONT_FAMILY
            if c_idx in active_cols:
                cell.fill.fore_color.rgb = C_CYAN
                run.text = "●"
                run.font.size = Pt(12)
                run.font.color.rgb = C_WHITE
            else:
                cell.fill.fore_color.rgb = C_WHITE
                run.text = "·"
                run.font.size = Pt(10)
                run.font.color.rgb = C_MUTED

    # Row 9: Team Effort Allocation
    c_eff = table.cell(9, 0)
    c_eff.fill.solid()
    c_eff.fill.fore_color.rgb = C_CALLOUT_BG
    c_eff.text = ""
    p = c_eff.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Team Responsibility"
    r.font.name = FONT_FAMILY
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = C_INDIGO

    # Member 1 merge (cols 1-3)
    table.cell(9, 1).merge(table.cell(9, 3))
    c_m1 = table.cell(9, 1)
    c_m1.fill.solid()
    c_m1.fill.fore_color.rgb = C_CRIMSON
    c_m1.text = ""
    p = c_m1.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Member 1: Testbed + Red-Teaming"
    r.font.name = FONT_FAMILY
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = C_WHITE

    # Member 2 merge (cols 4-6)
    table.cell(9, 4).merge(table.cell(9, 6))
    c_m2 = table.cell(9, 4)
    c_m2.fill.solid()
    c_m2.fill.fore_color.rgb = C_AMBER
    c_m2.text = ""
    p = c_m2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Member 2: System-1 Gate + Drift Encoder"
    r.font.name = FONT_FAMILY
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = C_WHITE

    # Member 3 merge (cols 7-8)
    table.cell(9, 7).merge(table.cell(9, 8))
    c_m3 = table.cell(9, 7)
    c_m3.fill.solid()
    c_m3.fill.fore_color.rgb = C_GREEN
    c_m3.text = ""
    p = c_m3.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Member 3: System-2 + HITL + Benchmarks"
    r.font.name = FONT_FAMILY
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = C_WHITE
