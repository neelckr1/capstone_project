"""
Slide 10: Conclusion & Thank You Builder
"""
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN
from deck_builder.config import (
    C_INDIGO, C_WHITE, C_LIGHT_INDIGO, C_CALLOUT_BG,
    C_CRIMSON, C_DARK, C_GREEN
)
from deck_builder.helpers import (
    clear_tf, add_run, add_textbox
)

def build_slide_10(slide):
    for shape in slide.shapes:
        if shape.name == "Rectangle 3":
            shape.fill.solid()
            shape.fill.fore_color.rgb = C_INDIGO
            tf = clear_tf(shape)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_run(p, "Thank You\n", bold=True, size=36, color=C_WHITE)
            add_run(p, "AegisMAS — Defending the Autonomous Frontier\n", bold=False, size=16, color=C_WHITE)
            add_run(p, "Questions & Panel Feedback Welcome", bold=False, italic=True, size=13, color=C_LIGHT_INDIGO)

    # 3 Metric Highlights above Rectangle 3 on Slide 10
    m_w = Emu(3200000)
    m_h = Emu(1300000)
    m_y = Emu(1600000)

    b1 = add_textbox(slide, Emu(1100000), m_y, m_w, m_h, "", bg_color=C_CALLOUT_BG, align=PP_ALIGN.CENTER)
    p_b1 = b1.text_frame.paragraphs[0]
    p_b1.alignment = PP_ALIGN.CENTER
    add_run(p_b1, "R₀ > 1  ➔  R₀ < 1\n", bold=True, size=22, color=C_CRIMSON)
    add_run(p_b1, "Viral Suppression Goal", bold=False, size=12, color=C_DARK)

    b2 = add_textbox(slide, Emu(4500000), m_y, m_w, m_h, "", bg_color=C_CALLOUT_BG, align=PP_ALIGN.CENTER)
    p_b2 = b2.text_frame.paragraphs[0]
    p_b2.alignment = PP_ALIGN.CENTER
    add_run(p_b2, "<50ms  Triage\n", bold=True, size=22, color=C_INDIGO)
    add_run(p_b2, "Non-Autoregressive Gate", bold=False, size=12, color=C_DARK)

    b3 = add_textbox(slide, Emu(7900000), m_y, m_w, m_h, "", bg_color=C_CALLOUT_BG, align=PP_ALIGN.CENTER)
    p_b3 = b3.text_frame.paragraphs[0]
    p_b3.alignment = PP_ALIGN.CENTER
    add_run(p_b3, ">75%  Cost Savings\n", bold=True, size=22, color=C_GREEN)
    add_run(p_b3, "Compute Efficiency", bold=False, size=12, color=C_DARK)

    # Bottom banner on Slide 10
    add_textbox(
        slide, Emu(1100000), Emu(6100000), Emu(10000000), Emu(450000),
        "UE24CS320A Capstone Project  •  PES University  •  2025–26",
        size=12, bold=True, color=C_WHITE, bg_color=C_INDIGO, align=PP_ALIGN.CENTER
    )
